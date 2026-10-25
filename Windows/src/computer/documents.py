"""
Document Text Extraction.

Extracts text from common file types so that Memory can compare
versions and the agent can read file content.
"""

import logging
import subprocess
import tempfile
from functools import lru_cache
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

# ═══════════════════════════════════════════════════════════
# Supported text-based extensions
# ═══════════════════════════════════════════════════════════

TEXT_EXTENSIONS = {
    # Programming
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".kt", ".go",
    ".rs", ".rb", ".php", ".c", ".h", ".cpp", ".hpp", ".cs",
    ".swift", ".lua", ".r", ".R", ".pl", ".scala", ".dart",
    # Web
    ".html", ".htm", ".css", ".scss", ".sass", ".less", ".vue", ".svelte",
    # Data / config
    ".json", ".yaml", ".yml", ".toml", ".xml", ".csv", ".tsv",
    ".ini", ".cfg", ".conf", ".env",
    # Docs
    ".md", ".rst", ".txt", ".tex", ".adoc",
    # Shell
    ".sh", ".bash", ".zsh", ".fish", ".bat", ".ps1",
    # Other
    ".sql", ".vim", ".dockerfile", ".gitignore", ".editorconfig",
    ".log",
}

# Names that count as text regardless of extension
TEXT_NAMES = {
    "Dockerfile", "Makefile", "Rakefile", "Vagrantfile",
    "Gemfile", "Procfile", "Justfile",
}


def extract_text(file_path: str) -> Optional[str]:
    """Extract text from *file_path*.

    Returns the text content, or ``None`` if the format is unsupported
    or the file cannot be read.
    """
    path = Path(file_path)
    if not path.exists() or not path.is_file():
        return None

    suffix = path.suffix.lower()

    if suffix in TEXT_EXTENSIONS or path.name in TEXT_NAMES:
        return _read_text(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    # Extensionless files — try as text
    if not suffix:
        return _read_text(path)

    return None


def extract_pages(file_path: str) -> list[tuple[Optional[int], str]]:
    """Extract text with PDF page numbers where the parser provides them."""
    path = Path(file_path)
    if not path.exists() or not path.is_file():
        return []
    if path.suffix.lower() == ".pdf":
        extracted = []
        try:
            import pymupdf
            with pymupdf.open(path) as document:
                for index, page in enumerate(document):
                    text = page.get_text().strip()
                    if not text:
                        text = _ocr_pdf_page(page)
                    if text.strip():
                        extracted.append((index + 1, text))
                if extracted:
                    return extracted
        except ImportError:
            pass
        except Exception as exc:
            logger.debug("PDF page extraction failed for %s: %s", path, exc)
        try:
            result = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, timeout=30)
            if result.returncode == 0:
                return [(index + 1, page) for index, page in enumerate(result.stdout.split("\f")) if page.strip()]
        except (FileNotFoundError, subprocess.TimeoutExpired):
            pass
        return []
    content = extract_text(str(path))
    return [(None, content)] if content and content.strip() else []


def _ocr_pdf_page(page) -> str:
    """OCR a scanned PDF page with the bundled RapidOCR engine."""
    try:
        import pymupdf
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
        with tempfile.NamedTemporaryFile(suffix=".png") as image_file:
            image_file.write(pixmap.tobytes("png"))
            image_file.flush()
            result = _rapid_ocr()(image_file.name)
        if not result:
            return ""
        recognized = getattr(result, "txts", None)
        if recognized is not None:
            return "\n".join(str(text) for text in recognized if text)
        return "\n".join(str(line[1]) for line in result if len(line) > 1)
    except Exception as exc:
        logger.debug("PDF OCR unavailable or failed: %s", exc)
        return ""


@lru_cache(maxsize=1)
def _rapid_ocr():
    from rapidocr import RapidOCR
    return RapidOCR()


def is_text_file(file_path: str) -> bool:
    """Quick check whether *file_path* is a known text format."""
    p = Path(file_path)
    return (
        p.suffix.lower() in TEXT_EXTENSIONS
        or p.name in TEXT_NAMES
        or p.suffix.lower() == ".docx"
    )


# ═══════════════════════════════════════════════════════════
# Extractors
# ═══════════════════════════════════════════════════════════


def _read_text(path: Path, max_bytes: int = 1_000_000) -> Optional[str]:
    try:
        if path.stat().st_size > max_bytes:
            logger.debug("File too large (%d B): %s", path.stat().st_size, path)
            return None
        return path.read_text(encoding="utf-8", errors="replace")
    except (OSError, UnicodeDecodeError) as exc:
        logger.debug("Cannot read %s: %s", path, exc)
        return None


def _extract_docx(path: Path) -> Optional[str]:
    try:
        from docx import Document  # type: ignore[import-untyped]

        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        paragraphs.extend(
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in doc.tables for row in table.rows
        )
        return "\n".join(paragraphs) if paragraphs else None
    except ImportError:
        logger.debug("python-docx not installed; skipping .docx")
        return None
    except Exception as exc:
        logger.debug("docx extraction failed for %s: %s", path, exc)
        return None


def _extract_pdf(path: Path) -> Optional[str]:
    # Prefer pdftotext (poppler-utils)
    try:
        result = subprocess.run(
            ["pdftotext", str(path), "-"],
            capture_output=True,
            text=True,
            timeout=30,
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    # Fallback: PyMuPDF
    try:
        import pymupdf

        doc = pymupdf.open(str(path))
        text = "\n".join(page.get_text() for page in doc)
        doc.close()
        return text if text.strip() else None
    except ImportError:
        logger.debug("No PDF extractor available (install poppler-utils or PyMuPDF)")
        return None
    except Exception as exc:
        logger.debug("PDF extraction failed for %s: %s", path, exc)
        return None
